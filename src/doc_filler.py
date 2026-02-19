"""
Módulo de preenchimento de modelos de contrato Word (DOCX).
Substitui placeholders {{...}} pelos dados extraídos dos PDFs.
"""

import os
import re
import copy
from docx import Document
from typing import Dict, Optional


def preencher_contrato(caminho_template: str, dados: Dict[str, str], caminho_saida: str) -> str:
    """
    Abre um template Word, substitui todos os placeholders pelos dados fornecidos,
    e salva o documento resultante no caminho de saída.
    
    Args:
        caminho_template: Caminho para o arquivo .docx template
        dados: Dicionário com os dados {PLACEHOLDER: valor}
        caminho_saida: Caminho para salvar o documento preenchido
    
    Returns:
        Caminho do arquivo gerado
    """
    doc = Document(caminho_template)
    
    # Substituir em parágrafos
    for paragrafo in doc.paragraphs:
        _substituir_placeholder_paragrafo(paragrafo, dados)
    
    # Substituir em tabelas
    for tabela in doc.tables:
        for linha in tabela.rows:
            for celula in linha.cells:
                for paragrafo in celula.paragraphs:
                    _substituir_placeholder_paragrafo(paragrafo, dados)
    
    # Substituir em cabeçalhos e rodapés
    for secao in doc.sections:
        for paragrafo in secao.header.paragraphs:
            _substituir_placeholder_paragrafo(paragrafo, dados)
        for paragrafo in secao.footer.paragraphs:
            _substituir_placeholder_paragrafo(paragrafo, dados)
    
    # Criar diretório de saída se não existir
    os.makedirs(os.path.dirname(caminho_saida) if os.path.dirname(caminho_saida) else ".", exist_ok=True)
    
    doc.save(caminho_saida)
    return caminho_saida


def _substituir_placeholder_paragrafo(paragrafo, dados: Dict[str, str]):
    """
    Substitui placeholders em um parágrafo.
    Mantém fonte/tamanho originais e deixa apenas os
    valores preenchidos em negrito.
    """
    # Verificar rapidamente se há placeholders no parágrafo
    texto_completo = paragrafo.text
    if "{{" not in texto_completo or "}}" not in texto_completo:
        return

    # Encontrar todos os placeholders no texto original (antes de substituir)
    matches = list(re.finditer(r'\{\{([^}]+)\}\}', texto_completo))
    if not matches:
        return

    runs = paragrafo.runs
    if not runs:
        return

    primeiro_run = runs[0]

    # Quebrar o texto em segmentos estáticos e de valor
    segmentos = []
    ultimo_fim = 0
    for m in matches:
        inicio, fim = m.span()
        # Texto estático antes do placeholder
        if inicio > ultimo_fim:
            segmentos.append(("static", texto_completo[ultimo_fim:inicio]))

        chave_bruta = m.group(1)
        chave = chave_bruta.strip()
        # Se não houver dado para a chave, mantemos o placeholder como texto estático
        if chave in dados:
            valor = dados[chave] if dados[chave] is not None else ""
            segmentos.append(("valor", valor))
        else:
            segmentos.append(("static", texto_completo[inicio:fim]))

        ultimo_fim = fim

    # Sobra de texto após o último placeholder
    if ultimo_fim < len(texto_completo):
        segmentos.append(("static", texto_completo[ultimo_fim:]))

    # Limpar texto de todos os runs existentes
    for run in runs:
        run.text = ""

    # Helper para copiar a formatação base (fonte/tamanho) do primeiro run
    def _configurar_run(run_destino, negrito: bool):
        # Manter estilo base
        try:
            run_destino.style = primeiro_run.style
        except Exception:
            pass

        if primeiro_run.font is not None and run_destino.font is not None:
            # Fonte e tamanho originais
            if primeiro_run.font.name:
                run_destino.font.name = primeiro_run.font.name
            if primeiro_run.font.size:
                run_destino.font.size = primeiro_run.font.size

        run_destino.bold = negrito

    # Criar novamente os runs com base nos segmentos
    run_atual = primeiro_run
    primeiro_segmento = True

    for tipo, texto in segmentos:
        if not texto:
            continue

        # Reusar o primeiro run apenas no primeiro segmento;
        # depois disso, criar novos runs.
        if primeiro_segmento:
            primeiro_segmento = False
        else:
            run_atual = paragrafo.add_run()

        if tipo == "valor":
            # Texto inserido pelo sistema: sempre em negrito
            _configurar_run(run_atual, negrito=True)
        else:
            # Texto fixo do contrato: nunca em negrito
            _configurar_run(run_atual, negrito=False)

        run_atual.text = texto


def gerar_contratos(dados_empregado: Dict[str, str], templates_selecionados: list, 
                     pasta_saida: str) -> list:
    """
    Gera contratos preenchidos para um empregado.
    
    Args:
        dados_empregado: Dicionário com os dados do empregado
        templates_selecionados: Lista de caminhos dos templates a preencher
        pasta_saida: Pasta onde salvar os documentos gerados
    
    Returns:
        Lista de caminhos dos arquivos gerados
    """
    nome_empregado = dados_empregado.get("NOME", "empregado").replace(" ", "_")
    arquivos_gerados = []
    
    for template_path in templates_selecionados:
        nome_template = os.path.splitext(os.path.basename(template_path))[0]
        nome_saida = f"{nome_template} - {dados_empregado.get('NOME', 'empregado')}.docx"
        caminho_saida = os.path.join(pasta_saida, nome_empregado, nome_saida)
        
        try:
            arquivo = preencher_contrato(template_path, dados_empregado, caminho_saida)
            arquivos_gerados.append(arquivo)
        except Exception as e:
            arquivos_gerados.append(f"ERRO: {template_path} - {str(e)}")
    
    return arquivos_gerados


def listar_templates(pasta_modelos: str = "Modelos") -> list:
    """Lista os templates disponíveis na pasta de modelos."""
    templates = []
    if os.path.exists(pasta_modelos):
        for arquivo in sorted(os.listdir(pasta_modelos)):
            if arquivo.endswith(".docx") and not arquivo.startswith("~$"):
                templates.append(os.path.join(pasta_modelos, arquivo))
    return templates
